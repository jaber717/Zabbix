"""Command line for the reporting suite."""
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from . import SUITE_VERSION
from . import delivery as dlv
from . import output as out
from .apiclient import ApiClient
from .config import ConfigError, REPORT_KEYS, load_all
from .model import fmt_kpi
from .periods import containing, previous_complete
from .plans import REPORTS, collect, resolve

MIME = {"pdf": "application/pdf", "json": "application/json",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def _base_loader():
    here = Path(__file__).resolve().parents[2] / "bin"
    if str(here) not in sys.path:
        sys.path.insert(0, str(here))
    from zabbix_daily_report import load_config          # the unchanged RC2 loader/validator
    return load_config


def parse_args(argv):
    p = argparse.ArgumentParser(prog="zabbix_report_suite.py", description="PDF / XLSX / JSON reports from Zabbix")
    p.add_argument("--config", help="report.json (RC2 configuration); required except for --selftest and --print-calendar")
    p.add_argument("--suite-config", help="suite.json (optional; defaults apply when absent)")
    p.add_argument("--report", default="all", help="all, or a comma list: daily, wan, infra, executive, incident, quality")
    p.add_argument("--cadence", choices=["daily", "weekly", "monthly"], help="run every report of this cadence (used by the timers)")
    p.add_argument("--selftest", action="store_true", help="render the bundled synthetic dataset to a temporary directory and verify PDF/XLSX/JSON (no Zabbix needed)")
    p.add_argument("--print-calendar", choices=["daily", "weekly", "monthly"], help="print the systemd OnCalendar value for a cadence from the suite configuration")
    p.add_argument("--date", help="any calendar date inside the wanted period (default: the last COMPLETE period)")
    p.add_argument("--formats", help="comma list of pdf,xlsx,json (default: suite formats)")
    p.add_argument("--output-directory")
    p.add_argument("--fixture", help="build from a saved dataset JSON instead of the Zabbix API")
    p.add_argument("--plan", action="store_true", help="print the planned periods and exit (no API access, no files)")
    p.add_argument("--dry-run", action="store_true", help="read from Zabbix and print KPIs; write nothing, send nothing")
    p.add_argument("--no-email", action="store_true", help="never send e-mail (overrides --send)")
    p.add_argument("--send", action="store_true", help="deliver by e-mail if delivery is enabled and permitted")
    p.add_argument("--resend", action="store_true", help="deliver again even if this exact report was already delivered")
    p.add_argument("--force", action="store_true", help="regenerate even when an identical complete run exists")
    p.add_argument("--now", help="ISO timestamp to use as the current time (testing)")
    return p.parse_args(argv)


def _now(a, tz):
    return datetime.fromisoformat(a.now).astimezone(tz) if a.now else datetime.now(tz)


def _period(kind, a, now, tz, week_start):
    if a.date:
        return containing(kind, date.fromisoformat(a.date), tz, week_start)
    return previous_complete(kind, now, tz, week_start)


def _render(doc, dataset, formats, include_dataset):
    from .render_json import render_json
    rendered = {}
    for fmt in formats:
        if fmt == "json":
            rendered[fmt] = render_json(doc, dataset if include_dataset else None)
        elif fmt == "pdf":
            from .render_pdf import render_pdf
            rendered[fmt] = render_pdf(doc)
        elif fmt == "xlsx":
            from .render_xlsx import render_xlsx
            rendered[fmt] = render_xlsx(doc)
    return rendered


def main(argv=None, api_factory=None, environ=None, smtp_factory=None, stdout=None):
    say = (lambda s: print(s, file=stdout or sys.stdout))
    a = parse_args(argv if argv is not None else sys.argv[1:])
    environ = os.environ if environ is None else environ
    if a.selftest:
        return selftest(say)
    if a.print_calendar:
        from .config import load_suite, on_calendar
        try:
            say(on_calendar(load_suite(a.suite_config), a.print_calendar))
            return 0
        except (ConfigError, ValueError, OSError, KeyError) as exc:
            print("RESULT=FAIL ERROR=configuration: %s" % exc, file=sys.stderr)
            return 2
    if not a.config:
        print("RESULT=FAIL ERROR=--config is required", file=sys.stderr)
        return 2
    try:
        cfg = load_all(a.config, a.suite_config, _base_loader())
        if a.cadence:
            keys = tuple(k for k in REPORT_KEYS if REPORTS[k]["kind"] == a.cadence)
        else:
            keys = REPORT_KEYS if a.report == "all" else tuple(resolve(x) for x in a.report.split(","))
        keys = [k for k in keys if cfg["suite"]["reports"].get(k, {}).get("enabled", True)]
    except (ConfigError, ValueError, OSError, KeyError) as exc:
        print("RESULT=FAIL ERROR=configuration: %s" % exc, file=sys.stderr)
        return 2
    suite = cfg["suite"]
    tz = ZoneInfo(cfg["report"]["timezone"])
    now = _now(a, tz)
    week_start = 6 if suite["week_start"] == "sunday" else 0
    formats = [f.strip() for f in a.formats.split(",")] if a.formats else list(suite["formats"])
    if not set(formats) <= {"pdf", "xlsx", "json"}:
        print("RESULT=FAIL ERROR=unknown format in --formats", file=sys.stderr)
        return 2
    base_dir = a.output_directory or suite["output"]["directory"]
    plan = [(k, _period(REPORTS[k]["kind"], a, now, tz, week_start)) for k in keys]
    if a.plan:
        for k, p in plan:
            say("PLAN report=%s kind=%s period=%s start=%s end=%s(exclusive) formats=%s" % (k, p.kind, p.label, p.start.isoformat(), p.end.isoformat(), ",".join(formats)))
        say("RESULT=PASS")
        return 0
    api = None
    fixture = json.loads(Path(a.fixture).read_text()) if a.fixture else None
    if fixture is None:
        api = api_factory(cfg["zabbix"]) if api_factory else ApiClient(cfg["zabbix"])
        try:
            api.authenticate()
        except Exception as exc:       # noqa: BLE001
            print("RESULT=FAIL ERROR=Zabbix authentication/connection: %s" % exc, file=sys.stderr)
            return 1
    failures, kept = 0, set()
    fm, dm = int(suite["output"]["file_mode"], 8), int(suite["output"]["dir_mode"], 8)
    from .reports import build
    for key, period in plan:
        try:
            if fixture is not None:
                ds = fixture["reports"][key] if "reports" in fixture else fixture
                if ds.get("report") != key:
                    raise ValueError("fixture does not contain report %s" % key)
            else:
                ds = collect(api, cfg, now, key, period, week_start)
            doc = build(ds, cfg)
            title = REPORTS[key]["title"]
            if a.dry_run:
                say("DRY_RUN report=%s period=%s dataset=%s warnings=%d api_calls=%d" % (key, ds["period"]["label"], ds["dataset_sha256"][:16], len(doc["warnings"]), ds["audit"]["api_calls"]))
                for k in doc["kpis"]:
                    say("  KPI %s = %s" % (k["id"], fmt_kpi(k)))
                for w in doc["warnings"]:
                    say("  WARNING %s" % w)
                continue
            d = out.run_dir(base_dir, REPORTS[key]["slug"], ds["period"]["label"], dm)
            kept.add(str(d))                          # a run is never expired by the invocation that just made it
            stem = "%s-%s" % (REPORTS[key]["slug"], ds["period"]["label"])
            if out.is_current(d, ds["dataset_sha256"], formats) and not a.force:
                status, manifest = "UNCHANGED", out.existing_manifest(d)
            else:
                rendered = _render(doc, ds, formats, suite["output"]["include_dataset_in_json"])
                manifest = out.write_outputs(d, stem, rendered, doc, ds["dataset_sha256"], fm,
                                             {"run_id": "%s-%s" % (ds["generated_at"], ds["dataset_sha256"][:12]), "generated_at": ds["generated_at"],
                                              "suite_version": SUITE_VERSION})
                status = "GENERATED"
            say("REPORT=%s PERIOD=%s STATUS=%s DATASET=%s WARNINGS=%d FILES=%s" % (key, ds["period"]["label"], status, ds["dataset_sha256"][:16], len(doc["warnings"]),
                                                                                  ",".join(f["name"] for f in manifest["files"])))
            _maybe_deliver(cfg, suite, a, environ, smtp_factory, d, doc, ds, manifest, title, say)
        except Exception as exc:       # noqa: BLE001 - one failing report must not stop the others
            failures += 1
            print("REPORT=%s STATUS=FAILED ERROR=%s: %s" % (key, type(exc).__name__, exc), file=sys.stderr)
    if not a.dry_run and failures < len(plan):
        for r in out.apply_retention(base_dir, suite["output"]["retention_days"], int(now.timestamp()), kept):
            say("RETENTION_REMOVED=%s" % r)
    say("RESULT=%s" % ("PASS" if not failures else "FAIL"))
    return 1 if failures else 0


def _maybe_deliver(cfg, suite, a, environ, smtp_factory, d, doc, ds, manifest, title, say):
    recipients, reason = dlv.decide_recipients(suite["delivery"], a.send, a.no_email, environ)
    if not recipients:
        if a.send or suite["delivery"]["enabled"]:
            say("DELIVERY=SKIPPED REASON=%s" % reason)
        return
    hashes = [f["sha256"] for f in manifest["files"] if f["format"] in ("pdf", "xlsx")]
    key = dlv.delivery_key(doc["report"], ds["period"]["label"], ds["dataset_sha256"], recipients, hashes)
    if dlv.already_delivered(d, key) and not a.resend:
        say("DELIVERY=ALREADY_DELIVERED RECIPIENTS=%d" % len(recipients))
        return
    attachments = []
    for f in manifest["files"]:
        if f["format"] in ("pdf", "xlsx"):
            attachments.append((f["name"], MIME[f["format"]], (Path(d) / f["name"]).read_bytes()))
    body = "%s\n%s\n\n%s\n\nGenerated %s. Dataset %s.\n" % (title, ds["period"]["label"], "\n".join("- " + x for x in doc["findings"]), ds["generated_at"], ds["dataset_sha256"][:16])
    if doc["warnings"]:
        body += "\nINCOMPLETE OR DEGRADED DATA:\n" + "\n".join("- " + w for w in doc["warnings"][:10]) + "\n"
    msg = dlv.build_message(cfg["delivery"]["sender"], recipients, "%s %s %s" % (suite["delivery"]["subject_prefix"], title, ds["period"]["label"]),
                            body, attachments, suite["delivery"]["max_attachment_mb"])
    kw = {"smtp_factory": smtp_factory} if smtp_factory else {}
    dlv.smtp_send(cfg["delivery"], msg, environ, **kw)
    dlv.record(d, key, {"recipients": len(recipients), "mode": suite["delivery"]["mode"], "at": ds["generated_at"]}, int(suite["output"]["file_mode"], 8))
    say("DELIVERY=SENT RECIPIENTS=%d MODE=%s" % (len(recipients), suite["delivery"]["mode"]))


def selftest(say):
    """Proves the installed dependencies and renderers work, using a bundled synthetic dataset."""
    import tempfile
    from .reports import build
    path = Path(__file__).resolve().parent / "data" / "selftest.json"
    fx = json.loads(path.read_text())
    cfg, ds = fx["cfg"], fx["dataset"]
    doc = build(ds, cfg)
    with tempfile.TemporaryDirectory(prefix="zrs-selftest-") as tmp:
        rendered = _render(doc, ds, ["json", "pdf", "xlsx"], True)
        d = Path(tmp)
        manifest = out.write_outputs(d, "selftest", rendered, doc, ds["dataset_sha256"], 0o600,
                                     {"run_id": "selftest", "generated_at": ds["generated_at"], "suite_version": SUITE_VERSION})
        checks = []
        pdf = (d / "selftest.pdf").read_bytes()
        checks.append(("PDF", pdf.startswith(b"%PDF-") and ("dataset-sha256:" + ds["dataset_sha256"]).encode() in pdf))
        from openpyxl import load_workbook
        wb = load_workbook(str(d / "selftest.xlsx"))
        checks.append(("XLSX", wb.sheetnames[0] == "Summary" and "Audit" in wb.sheetnames))
        j = json.loads((d / "selftest.json").read_text())
        checks.append(("JSON", j["dataset_sha256"] == ds["dataset_sha256"] and [k["value"] for k in j["kpis"]] == [k["value"] for k in doc["kpis"]]))
        checks.append(("MANIFEST", manifest["status"] == "COMPLETE" and len(manifest["files"]) == 3))
    for name, ok in checks:
        say("SELFTEST_%s=%s" % (name, "PASS" if ok else "FAIL"))
    say("SELFTEST_VERSIONS=%s" % json.dumps(out.versions(), sort_keys=True))
    ok = all(c[1] for c in checks)
    say("RESULT=%s" % ("PASS" if ok else "FAIL"))
    return 0 if ok else 1
