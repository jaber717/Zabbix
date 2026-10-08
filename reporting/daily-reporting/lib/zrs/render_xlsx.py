"""XLSX renderer (openpyxl). Values only: no formulas, no macros, typed numeric/date cells."""
from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .model import fmt_kpi

NUMFMT = {"int": "#,##0", "float1": "0.0", "float2": "0.00", "pct": '0.00"%"', "bps": "#,##0", "seconds": "#,##0", "datetime": "yyyy-mm-dd hh:mm"}
SUFFIX = {"bps": " (bps)", "seconds": " (s)", "pct": " (%)"}
RED, AMBER, GREEN = PatternFill("solid", bgColor="F8D0CC", fgColor="F8D0CC"), PatternFill("solid", bgColor="FFE9B3", fgColor="FFE9B3"), PatternFill("solid", bgColor="CDEBD7", fgColor="CDEBD7")
STATUS_FILL = {"DOWN": RED, "OPEN": RED, "Disaster": RED, "High": RED, "LOW": RED, "NONE": RED, "Average": AMBER, "Warning": AMBER,
               "MEDIUM": AMBER, "UNKNOWN": AMBER, "UP": GREEN, "RESOLVED": GREEN, "HIGH": GREEN}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")


def _sheet_name(title, used):
    base = re.sub(r"[\[\]\*\?/\\:]", "-", title)[:28] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        name = "%s~%d" % (base[:25], n)
        n += 1
    used.add(name.lower())
    return name


def _put(ws, row, column, value, ctype=None):
    """Write one cell. Strings are always stored as text: a leading = + - @ can never become a formula."""
    cell = ws.cell(row=row, column=column)
    if value is None:
        return cell
    if ctype == "datetime":
        try:
            dt = datetime.fromisoformat(value).replace(tzinfo=None)
            cell.value = dt
            cell.number_format = NUMFMT["datetime"]
            return cell
        except (TypeError, ValueError):
            pass
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
        if value[:1] in ("=", "+", "-", "@"):
            cell.quotePrefix = True
        return cell
    cell.value = value
    if ctype in NUMFMT:
        cell.number_format = NUMFMT[ctype]
    return cell


def _write_table(wb, used, t, tz_note=""):
    ws = wb.create_sheet(_sheet_name(t["title"], used))
    cols = t["columns"]
    for ci, c in enumerate(cols, start=1):
        h = ws.cell(row=1, column=ci, value=c["label"] + SUFFIX.get(c["type"], ""))
        h.font, h.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        h.alignment = Alignment(wrap_text=True, vertical="top")
    for ri, r in enumerate(t["rows"], start=2):
        for ci, c in enumerate(cols, start=1):
            _put(ws, ri, ci, r.get(c["key"]), c["type"])
    n = max(len(t["rows"]), 1)
    ref = "A1:%s%d" % (get_column_letter(len(cols)), n + 1)
    if t["rows"]:
        tab = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", "T_" + t["id"]), ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
    ws.freeze_panes = "B2" if len(cols) > 3 else "A2"
    for ci, c in enumerate(cols, start=1):
        width = max([len(str(c["label"])) + 2] + [len(str(r.get(c["key"]) if r.get(c["key"]) is not None else "")) for r in t["rows"][:200]])
        ws.column_dimensions[get_column_letter(ci)].width = min(max(width, 8), 60)
        cf = c.get("cf")
        rng = "%s2:%s%d" % (get_column_letter(ci), get_column_letter(ci), n + 1)
        if cf and t["rows"]:
            if cf["mode"] == "high":
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[repr(float(cf["bad"]))], fill=RED))
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[repr(float(cf["warn"]))], fill=AMBER))
            elif cf["mode"] == "low":
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[repr(float(cf["bad"]))], fill=RED))
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[repr(float(cf["warn"]))], fill=AMBER))
            elif cf["mode"] == "status":
                for word, fill in STATUS_FILL.items():
                    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % word], fill=fill))
    return ws


def render_xlsx(doc):
    wb = Workbook()
    wb.properties.creator = "Zabbix Reporting Suite"
    wb.properties.title = "%s %s" % (doc["title"], doc["period"]["label"])
    try:
        wb.properties.created = wb.properties.modified = datetime.fromisoformat(doc["generated_at"]).replace(tzinfo=None)
    except ValueError:
        pass
    used = {"summary", "charts", "data dictionary", "audit"}
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = doc["title"]
    ws["A1"].font = Font(bold=True, size=16, color="1F3A5F")
    p = doc["period"]
    meta = [("Period", "%s  (%s to %s, half-open)" % (p["label"], p["start"], p["end"])), ("Timezone", doc["timezone"]),
            ("Generated", doc["generated_at"]), ("Source", doc["source"]["zabbix"]), ("Dataset SHA-256", doc["dataset_sha256"])]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        _put(ws, i, 2, v, "text")
    r0 = 3 + len(meta) + 1
    for ci, h in enumerate(["KPI", "Value", "Unit", "Display", "Status", "Note", "Definition"], start=1):
        c = ws.cell(row=r0, column=ci, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for i, k in enumerate(doc["kpis"], start=r0 + 1):
        ws.cell(row=i, column=1, value=k["label"])
        _put(ws, i, 2, k["value"], "float2" if isinstance(k["value"], float) else "int")
        ws.cell(row=i, column=3, value=k["unit"])
        ws.cell(row=i, column=4, value=fmt_kpi(k))
        ws.cell(row=i, column=5, value=k["status"])
        _put(ws, i, 6, k["note"], "text")
        ws.cell(row=i, column=7, value=k["formula"])
    r = r0 + len(doc["kpis"]) + 2
    ws.cell(row=r, column=1, value="Key findings").font = Font(bold=True)
    for j, f in enumerate(doc["findings"], start=r + 1):
        _put(ws, j, 1, f, "text")
    r = r + len(doc["findings"]) + 2
    if doc["warnings"]:
        ws.cell(row=r, column=1, value="Incomplete or degraded data").font = Font(bold=True, color="9C0006")
        for j, wv in enumerate(doc["warnings"], start=r + 1):
            _put(ws, j, 1, wv, "text")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["F"].width = 34, 22, 60
    ws.freeze_panes = "A3"
    sheets = {}
    for sec in doc["sections"]:
        for tid in sec["tables"]:
            t = doc["tables"][tid]
            sheets[tid] = (_write_table(wb, used, t), t)
    if doc["charts"]:
        cs = wb.create_sheet("Charts")
        top = 1
        for c in doc["charts"]:
            if c["table"] not in sheets or not sheets[c["table"]][1]["rows"]:
                continue
            tws, t = sheets[c["table"]]
            keys = [x["key"] for x in t["columns"]]
            rows_n = min(len(t["rows"]), c["max_items"])
            ch = LineChart() if c["type"] == "line" else BarChart()
            if c["type"] == "hbar":
                ch.type = "bar"
            ch.title, ch.height, ch.width = c["title"], 8, 20
            for s in c["series"]:
                ci = keys.index(s["key"]) + 1
                ref = Reference(tws, min_col=ci, min_row=1, max_row=rows_n + 1)
                ch.add_data(ref, titles_from_data=True)
            ch.set_categories(Reference(tws, min_col=keys.index(c["x"]) + 1, min_row=2, max_row=rows_n + 1))
            cs.add_chart(ch, "A%d" % top)
            top += 18
    dws = wb.create_sheet("Data Dictionary")
    for ci, h in enumerate(["Term", "Definition"], start=1):
        c = dws.cell(row=1, column=ci, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for i, d in enumerate(doc["dictionary"], start=2):
        _put(dws, i, 1, d["term"], "text")
        c = _put(dws, i, 2, d["definition"], "text")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    dws.column_dimensions["A"].width, dws.column_dimensions["B"].width = 44, 120
    dws.freeze_panes = "A2"
    aws = wb.create_sheet("Audit")
    a = doc["audit"]
    rows = [("Dataset SHA-256", a["dataset_sha256"]), ("Generated", a["generated_at"]), ("Collector version", a["collector_version"]),
            ("Period start", a["period"]["start"]), ("Period end (exclusive)", a["period"]["end"]), ("Timezone", a["period"]["timezone"]),
            ("API calls", a["api_calls"]), ("Truncated requests", len(a["truncated"])), ("Failed requests", len(a["failed"]))]
    rows += [("API calls: " + k, v) for k, v in a["api_by_method"].items()]
    rows += [("Truncated: " + t["what"], t["detail"]) for t in a["truncated"]] + [("Failed: " + f["what"], f["error"]) for f in a["failed"]]
    rows += [("Note", n) for n in a["data_notes"]] + [("Scope", str(a["scope"])), ("Limits", str(a["limits"]))]
    aws.cell(row=1, column=1, value="Item").font = Font(bold=True)
    aws.cell(row=1, column=2, value="Value").font = Font(bold=True)
    for i, (k, v) in enumerate(rows, start=2):
        _put(aws, i, 1, k, "text")
        _put(aws, i, 2, v, "text" if isinstance(v, str) else "int")
    aws.column_dimensions["A"].width, aws.column_dimensions["B"].width = 36, 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
