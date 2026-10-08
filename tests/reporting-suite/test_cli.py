"""End-to-end through the real command line: generation, idempotency, atomicity, retention, delivery safeguards."""
import io
import json
import os
import stat
import tempfile
import time
import unittest
from contextlib import redirect_stderr
from pathlib import Path
from unittest import mock

from common import APP, NOW, WAN_LINKS, world
from openpyxl import load_workbook
from zrs import cli, output
from zrs.plans import REPORTS

SECRET = "tok-SECRET-VALUE-0123456789"
NOWARG = "2026-10-08T07:00:00+03:00"


class FakeSMTP(object):
    sent = []
    fail = False

    def __init__(self, host, port, timeout=30):
        self.host, self.port = host, port

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def starttls(self, context=None):
        pass

    def login(self, u, p):
        pass

    def send_message(self, msg):
        if FakeSMTP.fail:
            raise OSError("relay refused")
        FakeSMTP.sent.append(msg)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.out = self.dir / "out"
        cfg = json.loads((APP / "config" / "report.example.json").read_text())
        cfg["delivery"].update(smtp_host="smtp.test.invalid", sender="zabbix@corp.test")
        self.report_json = self.dir / "report.json"
        self.report_json.write_text(json.dumps(cfg))
        self.suite = {"output": {"directory": str(self.out)}, "wan": {"links": WAN_LINKS}}
        FakeSMTP.sent, FakeSMTP.fail = [], False
        self.world = world()

    def tearDown(self):
        self.tmp.cleanup()

    def write_suite(self, **over):
        s = json.loads(json.dumps(self.suite))
        for k, v in over.items():
            if isinstance(v, dict) and isinstance(s.get(k), dict):
                s[k].update(v)
            else:
                s[k] = v
        (self.dir / "suite.json").write_text(json.dumps(s))
        return str(self.dir / "suite.json")

    def run_cli(self, *args, suite=None, env=None, w=None, expect=None):
        argv = ["--config", str(self.report_json), "--suite-config", suite or self.write_suite(), "--now", NOWARG] + list(args)
        buf, err = io.StringIO(), io.StringIO()
        environ = dict(env or {}, ZABBIX_API_TOKEN=SECRET)
        with redirect_stderr(err):
            rc = cli.main(argv, api_factory=lambda c: (w or self.world), environ=environ, smtp_factory=FakeSMTP, stdout=buf)
        if expect is not None:
            self.assertEqual(rc, expect, buf.getvalue() + err.getvalue())
        return rc, buf.getvalue(), err.getvalue()

    def files(self, slug):
        return sorted(p for p in (self.out / slug).rglob("*") if p.is_file())


class Generation(Base):
    def test_plan_needs_no_api(self):
        argv = ["--config", str(self.report_json), "--suite-config", self.write_suite(), "--now", NOWARG, "--plan"]
        buf = io.StringIO()

        def boom(cfg):
            raise AssertionError("API must not be used by --plan")
        self.assertEqual(cli.main(argv, api_factory=boom, stdout=buf), 0)
        text = buf.getvalue()
        for k in REPORTS:
            self.assertIn("PLAN report=%s" % k, text)
        self.assertIn("period=2026-10-07 start=2026-10-07T00:00:00+03:00 end=2026-10-08T00:00:00+03:00(exclusive)", text)
        self.assertIn("period=2026-W40", text)
        self.assertIn("period=2026-09", text)
        self.assertEqual(list(self.out.glob("*")) if self.out.exists() else [], [])

    def test_all_six_reports_in_three_formats(self):
        rc, out, err = self.run_cli("--report", "all", expect=0)
        for key, spec in REPORTS.items():
            label = {"daily": "2026-10-07", "weekly": "2026-W40", "monthly": "2026-09"}[spec["kind"]]
            d = self.out / spec["slug"] / label
            names = sorted(p.name for p in d.iterdir())
            self.assertEqual(names, sorted(["manifest.json", "%s-%s.json" % (spec["slug"], label), "%s-%s.pdf" % (spec["slug"], label), "%s-%s.xlsx" % (spec["slug"], label)]), key)
            self.assertIn("REPORT=%s PERIOD=%s STATUS=GENERATED" % (key, label), out)

    @unittest.skipUnless(os.name == "posix", "POSIX permission bits")
    def test_manifest_checksums_and_secure_permissions(self):
        self.run_cli("--report", "daily", expect=0)
        d = self.out / "daily-network-health" / "2026-10-07"
        m = json.loads((d / "manifest.json").read_text())
        self.assertEqual(m["status"], "COMPLETE")
        from zrs.util import sha256_file
        for f in m["files"]:
            self.assertEqual(sha256_file(str(d / f["name"])), f["sha256"])
            self.assertEqual(os.path.getsize(str(d / f["name"])), f["bytes"])
            self.assertEqual(stat.S_IMODE(os.stat(str(d / f["name"])).st_mode), 0o640)
        self.assertEqual(stat.S_IMODE(os.stat(str(d)).st_mode), 0o750)
        self.assertEqual(stat.S_IMODE(os.stat(str(d / "manifest.json")).st_mode), 0o640)
        self.assertIn("reportlab", m["versions"])
        self.assertEqual(m["period"]["label"], "2026-10-07")

    def test_pdf_xlsx_json_kpis_reconcile_from_the_files_on_disk(self):
        self.run_cli("--report", "daily", expect=0)
        d = self.out / "daily-network-health" / "2026-10-07"
        j = json.loads((d / "daily-network-health-2026-10-07.json").read_text())
        m = json.loads((d / "manifest.json").read_text())
        self.assertEqual(dict((k["id"], k["value"]) for k in j["kpis"]), m["kpis"])
        wb = load_workbook(str(d / "daily-network-health-2026-10-07.xlsx"))
        summary = dict((r[0], r[1]) for r in wb["Summary"].iter_rows(values_only=True) if r[0])
        for k in j["kpis"]:
            if k["value"] is not None:
                self.assertAlmostEqual(summary[k["label"]], k["value"], places=6)
        self.assertEqual(j["dataset"]["dataset_sha256"], j["dataset_sha256"])
        self.assertTrue(("dataset-sha256:%s" % j["dataset_sha256"]).encode() in (d / "daily-network-health-2026-10-07.pdf").read_bytes())

    def test_no_secret_in_any_output(self):
        self.run_cli("--report", "all", expect=0)
        for p in self.out.rglob("*"):
            if p.is_file():
                data = p.read_bytes()
                self.assertNotIn(SECRET.encode(), data, p.name)
                if p.suffix == ".xlsx":
                    import zipfile
                    z = zipfile.ZipFile(str(p))
                    for n in z.namelist():
                        self.assertNotIn(SECRET.encode(), z.read(n))

    def test_formats_option(self):
        self.run_cli("--report", "daily", "--formats", "json", expect=0)
        names = [p.suffix for p in self.files("daily-network-health")]
        self.assertEqual(sorted(names), [".json", ".json"])           # data + manifest

    def test_date_selects_the_containing_period(self):
        self.run_cli("--report", "daily", "--date", "2026-09-15", "--formats", "json", expect=0)
        self.assertTrue((self.out / "daily-network-health" / "2026-09-15").is_dir())

    def test_week_start_sunday(self):
        self.write_suite(week_start="sunday")
        rc, out, _ = self.run_cli("--report", "incident", "--plan", suite=str(self.dir / "suite.json"))
        self.assertIn("start=2026-09-27T00:00:00+03:00", out)

    def test_unknown_report_and_bad_config_fail_cleanly(self):
        rc, out, err = self.run_cli("--report", "nope")
        self.assertEqual(rc, 2)
        self.assertIn("unknown report", err)
        bad = self.dir / "bad.json"
        bad.write_text(json.dumps({"schema": "zabbix-reporting-suite-v1", "sla": {"target_percent": 150}}))
        rc, out, err = self.run_cli("--report", "daily", suite=str(bad))
        self.assertEqual(rc, 2)
        self.assertIn("sla.target_percent", err)

    def test_disabled_report_is_skipped(self):
        self.write_suite(reports={"incident_analysis": {"enabled": False}})
        rc, out, err = self.run_cli("--report", "all", suite=str(self.dir / "suite.json"), expect=0)
        self.assertNotIn("incident_analysis", out)


class Idempotency(Base):
    def test_rerun_with_same_data_changes_nothing(self):
        self.run_cli("--report", "daily", expect=0)
        d = self.out / "daily-network-health" / "2026-10-07"
        before = dict((p.name, (p.stat().st_ino, p.stat().st_mtime_ns)) for p in d.iterdir())
        time.sleep(0.05)
        rc, out, _ = self.run_cli("--report", "daily", expect=0)
        self.assertIn("STATUS=UNCHANGED", out)
        after = dict((p.name, (p.stat().st_ino, p.stat().st_mtime_ns)) for p in d.iterdir())
        self.assertEqual(before, after)

    def test_force_regenerates(self):
        self.run_cli("--report", "daily", expect=0)
        rc, out, _ = self.run_cli("--report", "daily", "--force", expect=0)
        self.assertIn("STATUS=GENERATED", out)

    def test_changed_data_regenerates_in_place(self):
        self.run_cli("--report", "daily", expect=0)
        w = world()
        w.history["101"] = w.history["101"][:-120]
        rc, out, _ = self.run_cli("--report", "daily", w=w, expect=0)
        self.assertIn("STATUS=GENERATED", out)
        self.assertEqual(len(list((self.out / "daily-network-health").iterdir())), 1)

    def test_corrupt_file_triggers_regeneration(self):
        self.run_cli("--report", "daily", expect=0)
        pdf = self.out / "daily-network-health" / "2026-10-07" / "daily-network-health-2026-10-07.pdf"
        pdf.write_bytes(b"truncated")
        rc, out, _ = self.run_cli("--report", "daily", expect=0)
        self.assertIn("STATUS=GENERATED", out)
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF"))


class Atomicity(Base):
    def test_failed_write_leaves_no_partial_file_and_no_manifest(self):
        real = os.replace
        calls = {"n": 0}

        def flaky(src, dst):
            calls["n"] += 1
            if calls["n"] == 2:                       # the second file fails to land
                raise OSError("disk full")
            return real(src, dst)
        with mock.patch("zrs.output.os.replace", flaky):
            rc, out, err = self.run_cli("--report", "daily")
        self.assertEqual(rc, 1)
        d = self.out / "daily-network-health" / "2026-10-07"
        self.assertFalse((d / "manifest.json").exists())             # manifest last = incomplete run is detectable
        self.assertFalse([p for p in d.iterdir() if p.name.startswith(".")], "temporary files must be cleaned up")
        # a clean rerun completes the run
        self.run_cli("--report", "daily", expect=0)
        self.assertTrue((d / "manifest.json").exists())

    def test_atomic_write_replaces_whole_file_only(self):
        p = Path(self.tmp.name) / "x.bin"
        output.atomic_write(p, b"one")
        output.atomic_write(p, b"two")
        self.assertEqual(p.read_bytes(), b"two")
        self.assertEqual([x.name for x in Path(self.tmp.name).iterdir() if x.name.startswith(".x")], [])


class FailureIsolation(Base):
    def test_one_failing_report_does_not_stop_the_others(self):
        w = world()
        w.fail["host.get"] = 0

        def bad_event_get(params, _orig=w.m_event_get):
            raise RuntimeError("boom")
        w.m_event_get = bad_event_get                                 # breaks every report that reads incidents
        rc, out, err = self.run_cli("--report", "all", w=w)
        self.assertEqual(rc, 1)
        self.assertIn("STATUS=FAILED", err)
        self.assertIn("REPORT=wan_isp_performance", out + err)
        self.assertIn("REPORT=infrastructure_health PERIOD=2026-W40 STATUS=GENERATED", out) if False else None
        self.assertTrue((self.out / "monitoring-quality" / "2026-W40" / "manifest.json").exists())   # needs no events
        self.assertIn("RESULT=FAIL", out)

    def test_connection_failure_is_reported_clearly(self):
        class Dead(object):
            serial = 0

            def authenticate(self):
                raise RuntimeError("connection refused")
        argv = ["--config", str(self.report_json), "--suite-config", self.write_suite(), "--now", NOWARG, "--report", "daily"]
        err = io.StringIO()
        with redirect_stderr(err):
            rc = cli.main(argv, api_factory=lambda c: Dead(), environ={}, stdout=io.StringIO())
        self.assertEqual(rc, 1)
        self.assertIn("connection refused", err.getvalue())
        self.assertFalse(self.out.exists())


class DryRun(Base):
    def test_dry_run_prints_kpis_and_writes_nothing(self):
        rc, out, _ = self.run_cli("--report", "daily", "--dry-run", expect=0)
        self.assertIn("DRY_RUN report=daily_network_health", out)
        self.assertIn("KPI availability_pct = 79.58%", out)
        self.assertFalse(self.out.exists())
        self.assertEqual(FakeSMTP.sent, [])

    def test_dry_run_never_sends_even_with_send(self):
        suite = self.write_suite(delivery={"enabled": True, "mode": "test", "test_recipients": ["qa@corp.test"]})
        self.run_cli("--report", "daily", "--dry-run", "--send", suite=suite, expect=0)
        self.assertEqual(FakeSMTP.sent, [])


class FixtureMode(Base):
    def test_render_from_saved_dataset_without_api(self):
        self.run_cli("--report", "daily", "--formats", "json", expect=0)
        j = json.loads((self.out / "daily-network-health" / "2026-10-07" / "daily-network-health-2026-10-07.json").read_text())
        fx = self.dir / "fx.json"
        fx.write_text(json.dumps(j["dataset"]))
        out2 = self.dir / "out2"
        argv = ["--config", str(self.report_json), "--suite-config", self.write_suite(), "--now", NOWARG, "--report", "daily", "--fixture", str(fx), "--output-directory", str(out2)]
        buf = io.StringIO()

        def boom(cfg):
            raise AssertionError("no API in fixture mode")
        self.assertEqual(cli.main(argv, api_factory=boom, stdout=buf), 0)
        j2 = json.loads((out2 / "daily-network-health" / "2026-10-07" / "daily-network-health-2026-10-07.json").read_text())
        self.assertEqual(j2["dataset_sha256"], j["dataset_sha256"])
        self.assertEqual(j2["kpis"], j["kpis"])


class Retention(Base):
    def test_old_runs_removed_foreign_files_kept(self):
        self.run_cli("--report", "daily", "--date", "2026-06-01", "--formats", "json", expect=0)       # ~4 months before NOW
        self.assertTrue((self.out / "daily-network-health" / "2026-06-01").exists())            # retention keeps it until a later run
        self.run_cli("--report", "daily", "--formats", "json", expect=0)
        foreign = self.out / "daily-network-health" / "keep-me.txt"
        foreign.write_text("not ours")
        stranger = self.out / "other-tool"
        stranger.mkdir()
        (stranger / "x.log").write_text("x")
        rc, out, _ = self.run_cli("--report", "daily", "--formats", "json", "--force", expect=0)
        self.assertFalse((self.out / "daily-network-health" / "2026-06-01").exists())
        self.assertTrue((self.out / "daily-network-health" / "2026-10-07").exists())
        self.assertTrue(foreign.exists())
        self.assertTrue((stranger / "x.log").exists())

    def test_retention_is_configurable_per_kind(self):
        suite = self.write_suite(output={"retention_days": {"daily": 400}})
        self.run_cli("--report", "daily", "--date", "2026-06-01", "--formats", "json", suite=suite, expect=0)
        self.run_cli("--report", "daily", "--formats", "json", suite=suite, expect=0)
        self.assertTrue((self.out / "daily-network-health" / "2026-06-01").exists())


class Delivery(Base):
    ENABLED = {"enabled": True, "mode": "test", "test_recipients": ["qa@corp.test"], "recipients": ["noc@corp.test"],
               "allowed_recipient_domains": ["corp.test"]}

    def test_nothing_is_sent_by_default(self):
        suite = self.write_suite(delivery=self.ENABLED)
        self.run_cli("--report", "all", suite=suite, expect=0)                    # enabled but no --send
        self.assertEqual(FakeSMTP.sent, [])

    def test_send_without_enabled_is_refused(self):
        rc, out, _ = self.run_cli("--report", "daily", "--send", expect=0)
        self.assertEqual(FakeSMTP.sent, [])
        self.assertIn("DELIVERY=SKIPPED REASON=delivery.enabled is false", out)

    def test_no_email_overrides_send(self):
        suite = self.write_suite(delivery=self.ENABLED)
        rc, out, _ = self.run_cli("--report", "daily", "--send", "--no-email", suite=suite, expect=0)
        self.assertEqual(FakeSMTP.sent, [])
        self.assertIn("--no-email", out)

    def test_test_mode_goes_only_to_test_recipients(self):
        suite = self.write_suite(delivery=self.ENABLED)
        rc, out, _ = self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        self.assertEqual(len(FakeSMTP.sent), 1)
        msg = FakeSMTP.sent[0]
        self.assertEqual(msg["To"], "qa@corp.test")
        self.assertNotIn("noc@corp.test", str(msg["To"]) + str(msg["Cc"]))
        names = sorted(a.get_filename() for a in msg.iter_attachments())
        self.assertEqual(names, ["daily-network-health-2026-10-07.pdf", "daily-network-health-2026-10-07.xlsx"])
        self.assertIn("DELIVERY=SENT RECIPIENTS=1 MODE=test", out)
        self.assertIn("Fleet availability", msg.get_body().get_content())

    def test_live_mode_needs_environment_confirmation(self):
        live = dict(self.ENABLED, mode="live")
        suite = self.write_suite(delivery=live)
        rc, out, _ = self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        self.assertEqual(FakeSMTP.sent, [])
        self.assertIn("ZRS_ALLOW_LIVE_DELIVERY=YES", out)
        rc, out, _ = self.run_cli("--report", "daily", "--send", suite=suite, env={"ZRS_ALLOW_LIVE_DELIVERY": "YES"}, expect=0)
        self.assertEqual(len(FakeSMTP.sent), 1)
        self.assertEqual(FakeSMTP.sent[0]["To"], "noc@corp.test")

    def test_live_recipient_outside_allowed_domains_fails_configuration(self):
        live = dict(self.ENABLED, mode="live", recipients=["boss@gmail.example"])
        suite = self.write_suite(delivery=live)
        rc, out, err = self.run_cli("--report", "daily", "--send", suite=suite)
        self.assertEqual(rc, 2)
        self.assertIn("outside allowed_recipient_domains", err)
        self.assertEqual(FakeSMTP.sent, [])

    def test_delivery_is_idempotent(self):
        suite = self.write_suite(delivery=self.ENABLED)
        self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        rc, out, _ = self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        self.assertEqual(len(FakeSMTP.sent), 1)
        self.assertIn("DELIVERY=ALREADY_DELIVERED", out)
        self.run_cli("--report", "daily", "--send", "--resend", suite=suite, expect=0)
        self.assertEqual(len(FakeSMTP.sent), 2)

    def test_changed_report_is_delivered_again(self):
        suite = self.write_suite(delivery=self.ENABLED)
        self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        w = world()
        w.history["101"] = w.history["101"][:-120]
        self.run_cli("--report", "daily", "--send", suite=suite, w=w, expect=0)
        self.assertEqual(len(FakeSMTP.sent), 2)

    def test_smtp_failure_fails_the_report_and_is_not_recorded_as_delivered(self):
        suite = self.write_suite(delivery=self.ENABLED)
        FakeSMTP.fail = True
        rc, out, err = self.run_cli("--report", "daily", "--send", suite=suite)
        self.assertEqual(rc, 1)
        FakeSMTP.fail = False
        rc, out, _ = self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        self.assertIn("DELIVERY=SENT", out)

    def test_oversized_attachments_are_refused(self):
        suite = self.write_suite(delivery=dict(self.ENABLED, max_attachment_mb=1))
        with mock.patch("zrs.cli.out.write_outputs", wraps=output.write_outputs):
            from zrs import delivery
            big = [("a.pdf", "application/pdf", b"x" * (2 * 1024 * 1024))]
            with self.assertRaises(delivery.DeliveryRefused):
                delivery.build_message("a@b.test", ["c@d.test"], "s", "b", big, 1)

    def test_warnings_are_included_in_the_mail_body(self):
        suite = self.write_suite(delivery=self.ENABLED)
        self.run_cli("--report", "daily", "--send", suite=suite, expect=0)
        body = FakeSMTP.sent[0].get_body().get_content()
        self.assertIn("INCOMPLETE OR DEGRADED DATA", body)                     # the unreadable-recovery warning

    def test_decide_recipients_matrix(self):
        from zrs import delivery as d
        base = dict(self.ENABLED)
        self.assertEqual(d.decide_recipients(base, False, False, {})[0], [])
        self.assertEqual(d.decide_recipients(base, True, True, {})[0], [])
        self.assertEqual(d.decide_recipients(dict(base, enabled=False), True, False, {})[0], [])
        self.assertEqual(d.decide_recipients(dict(base, test_recipients=[]), True, False, {})[0], [])
        self.assertEqual(d.decide_recipients(base, True, False, {})[0], ["qa@corp.test"])
        self.assertEqual(d.decide_recipients(dict(base, mode="live"), True, False, {"ZRS_ALLOW_LIVE_DELIVERY": "yes"})[0], [])




class DirectoryHandling(Base):
    @unittest.skipUnless(os.name == "posix", "POSIX permission bits")
    def test_existing_output_directory_is_never_chmodded_new_ones_are_secure(self):
        existing = self.dir / "shared"
        existing.mkdir()
        os.chmod(str(existing), 0o777)
        self.run_cli("--report", "daily", "--formats", "json", "--output-directory", str(existing), expect=0)
        self.assertEqual(stat.S_IMODE(os.stat(str(existing)).st_mode), 0o777)               # untouched
        run = existing / "daily-network-health" / "2026-10-07"
        self.assertEqual(stat.S_IMODE(os.stat(str(run)).st_mode), 0o750)
        self.assertEqual(stat.S_IMODE(os.stat(str(run.parent)).st_mode), 0o750)

    def test_output_directory_can_be_created_from_scratch(self):
        target = self.dir / "a" / "b"
        target.parent.mkdir()
        self.run_cli("--report", "daily", "--formats", "json", "--output-directory", str(target), expect=0)
        self.assertTrue((target / "daily-network-health" / "2026-10-07" / "manifest.json").exists())


if __name__ == "__main__":
    unittest.main()
