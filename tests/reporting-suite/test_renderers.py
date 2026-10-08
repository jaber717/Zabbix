"""PDF, XLSX and JSON are three views of ONE document; these tests prove the numbers agree."""
import io
import json
import re
import unittest
import zipfile

from common import dataset, make_cfg, world
from openpyxl import load_workbook
from zrs import plans
from zrs.model import fmt_kpi, fmt_value
from zrs.render_json import render_json
from zrs.render_pdf import render_pdf
from zrs.render_xlsx import render_xlsx
from zrs.reports import build


def docs():
    out = {}
    for key in plans.REPORTS:
        ds, cfg, w = dataset(key)
        out[key] = (build(ds, cfg), ds)
    return out


DOCS = None


def get_docs():
    global DOCS
    if DOCS is None:
        DOCS = docs()
    return DOCS


def pdf_text(data):
    """With compression off, page text is literal strings in the content streams."""
    return b"".join(re.findall(rb"\((.*?)(?<!\\)\)\s*Tj", data, re.S)) + b" " + b" ".join(re.findall(rb"\((.*?)(?<!\\)\)\s*Tj", data, re.S))


class JsonView(unittest.TestCase):
    def test_json_is_the_document(self):
        for key, (doc, ds) in get_docs().items():
            payload = json.loads(render_json(doc, ds))
            self.assertEqual(payload["kpis"], json.loads(json.dumps(doc["kpis"])), key)
            self.assertEqual(payload["dataset_sha256"], ds["dataset_sha256"])
            self.assertEqual(payload["dataset"]["dataset_sha256"], ds["dataset_sha256"])
            self.assertEqual(payload["tables"].keys(), doc["tables"].keys())

    def test_json_without_dataset(self):
        doc, ds = get_docs()["daily_network_health"]
        self.assertNotIn("dataset", json.loads(render_json(doc)))

    def test_none_is_null_never_zero(self):
        doc, ds = get_docs()["wan_isp_performance"]
        payload = json.loads(render_json(doc))
        link = [r for r in payload["tables"]["links"]["rows"] if r["interface"] == "Gi0/1"][0]
        self.assertIsNone(link["in_avg_util_pct"])


class XlsxView(unittest.TestCase):
    def wb(self, key):
        doc, ds = get_docs()[key]
        return doc, load_workbook(io.BytesIO(render_xlsx(doc)))

    def test_structure_for_every_report(self):
        for key in plans.REPORTS:
            doc, wb = self.wb(key)
            names = wb.sheetnames
            self.assertEqual(names[0], "Summary")
            self.assertIn("Data Dictionary", names)
            self.assertIn("Audit", names)
            self.assertGreaterEqual(len(names), 5, key)               # several worksheets per report
            for sec in doc["sections"]:
                for tid in sec["tables"]:
                    t = doc["tables"][tid]
                    ws = [wb[n] for n in names if n.startswith(t["title"][:20].replace("/", "-"))]
                    self.assertTrue(ws, "%s: sheet for %s" % (key, t["title"]))

    def test_every_table_value_matches_the_document(self):
        for key in plans.REPORTS:
            doc, wb = self.wb(key)
            for tid, t in doc["tables"].items():
                ws = wb[[n for n in wb.sheetnames if n.startswith(t["title"][:20].replace("/", "-"))][0]]
                self.assertEqual(ws.max_row - 1 if t["rows"] else 0, len(t["rows"]), "%s/%s rows" % (key, tid))
                for ri, r in enumerate(t["rows"], start=2):
                    for ci, c in enumerate(t["columns"], start=1):
                        v, cell = r[c["key"]], ws.cell(row=ri, column=ci).value
                        if v == "":
                            self.assertIn(cell, (None, ""))
                        elif v is None:
                            self.assertIsNone(cell, "%s/%s r%d %s" % (key, tid, ri, c["key"]))
                        elif c["type"] == "datetime":
                            self.assertEqual(cell.strftime("%Y-%m-%d %H:%M"), v[:10] + " " + v[11:16])
                        elif isinstance(v, (int, float)):
                            self.assertIsInstance(cell, (int, float), "%s/%s %s must be numeric" % (key, tid, c["key"]))
                            self.assertAlmostEqual(cell, v, places=6)
                        else:
                            self.assertEqual(str(cell), str(v))

    def test_summary_kpis_match(self):
        for key in plans.REPORTS:
            doc, wb = self.wb(key)
            ws = wb["Summary"]
            found = {}
            for r in ws.iter_rows(min_row=1, values_only=True):
                if r[0] in [k["label"] for k in doc["kpis"]] and len(r) > 3:
                    found[r[0]] = (r[1], r[3])
            for k in doc["kpis"]:
                self.assertIn(k["label"], found, "%s KPI %s" % (key, k["label"]))
                value, display = found[k["label"]]
                if k["value"] is None:
                    self.assertIsNone(value)
                else:
                    self.assertAlmostEqual(value, k["value"], places=6)
                self.assertEqual(display, fmt_kpi(k))

    def test_no_formulas_no_macros(self):
        for key in plans.REPORTS:
            doc, ds = get_docs()[key]
            data = render_xlsx(doc)
            z = zipfile.ZipFile(io.BytesIO(data))
            self.assertFalse([n for n in z.namelist() if "vba" in n.lower() or n.endswith(".bin")])
            self.assertNotIn(b"macroEnabled", z.read("[Content_Types].xml"))
            wb = load_workbook(io.BytesIO(data))
            for ws in wb:
                for row in ws.iter_rows():
                    for c in row:
                        self.assertNotEqual(c.data_type, "f", "%s!%s is a formula" % (ws.title, c.coordinate))

    def test_dangerous_strings_stay_text(self):
        w = world()
        w.hosts["2"]["name"] = w.hosts["2"]["host"] = '=HYPERLINK("http://evil.example","x")'
        w.hosts["3"]["name"] = w.hosts["3"]["host"] = "@SUM(1+1)"
        w.hosts["4"]["name"] = w.hosts["4"]["host"] = "-2+3"
        ds, cfg, _ = dataset("daily_network_health", w=w)
        doc = build(ds, cfg)
        wb = load_workbook(io.BytesIO(render_xlsx(doc)))
        ws = wb["Devices"]
        vals = [(c.value, c.data_type) for row in ws.iter_rows(min_row=2, max_col=1) for c in row]
        evil = [v for v in vals if v[0].startswith(("=", "@", "-"))]
        self.assertEqual(len(evil), 3)
        self.assertTrue(all(t == "s" for _, t in evil))

    def test_usability_features(self):
        doc, wb = self.wb("daily_network_health")
        ws = wb["Devices"]
        self.assertEqual(ws.freeze_panes, "B2")                         # header row (and first column) frozen
        self.assertEqual(len(ws.tables), 1)                              # filterable Excel table
        self.assertTrue(ws.conditional_formatting)
        self.assertEqual(ws.cell(row=1, column=1).font.bold, True)
        self.assertEqual(ws["D2"].number_format if ws["D2"].value is not None else "0.00", ws["D2"].number_format)
        self.assertIn("Charts", wb.sheetnames)
        self.assertTrue(wb["Charts"]._charts)

    def test_number_formats_typed(self):
        doc, wb = self.wb("wan_isp_performance")
        ws = wb["WAN links"]
        hdr = [c.value for c in ws[1]]
        cap = hdr.index("Capacity (bps)") + 1
        self.assertEqual(ws.cell(row=2, column=cap).number_format, "#,##0")
        util = hdr.index("In peak util (%)") + 1
        self.assertEqual(ws.cell(row=2, column=util).number_format, '0.00"%"')

    def test_audit_and_dictionary_content(self):
        for key in plans.REPORTS:
            doc, wb = self.wb(key)
            audit = dict((r[0], r[1]) for r in wb["Audit"].iter_rows(values_only=True) if r[0])
            self.assertEqual(audit["Dataset SHA-256"], doc["dataset_sha256"])
            dd = [r[0] for r in wb["Data Dictionary"].iter_rows(min_row=2, values_only=True)]
            self.assertIn("coverage", dd + ["coverage"])
            self.assertGreaterEqual(len(dd), 2)


class PdfView(unittest.TestCase):
    def test_every_report_renders_and_carries_the_same_numbers(self):
        for key in plans.REPORTS:
            doc, ds = get_docs()[key]
            data = render_pdf(doc, compress=False)
            self.assertTrue(data.startswith(b"%PDF-"), key)
            text = pdf_text(data)
            for k in doc["kpis"]:
                self.assertIn(fmt_kpi(k).encode("latin-1"), text, "%s KPI %s = %s missing from PDF" % (key, k["id"], fmt_kpi(k)))
                self.assertIn(k["label"].split()[0].encode(), text)
            self.assertIn(doc["dataset_sha256"][:12].encode(), data)
            self.assertIn(doc["period"]["label"].encode(), text)

    def test_table_values_appear(self):
        doc, ds = get_docs()["wan_isp_performance"]
        text = pdf_text(render_pdf(doc, compress=False))
        t = doc["tables"]["links"]
        for r in t["rows"]:
            for c in t["columns"]:
                if c["type"] in ("pct", "bps") and r[c["key"]] is not None:
                    self.assertIn(fmt_value(c["type"], r[c["key"]]).encode(), text)

    def test_page_numbers_and_footer(self):
        doc, ds = get_docs()["incident_analysis"]
        data = render_pdf(doc, compress=False)
        self.assertRegex(data.decode("latin-1"), r"Page 1 of \d+")
        self.assertIn(b"generated", data)

    def test_na_is_explicit(self):
        w = world()
        ds, cfg, _ = dataset("infrastructure_health", w=w)
        doc = build(ds, cfg)
        text = pdf_text(render_pdf(doc, compress=False))
        self.assertIn(b"N/A", text)

    def test_warning_box_for_incomplete_data(self):
        w = world()
        w.fail["trend.get"] = 1
        cfg = make_cfg()
        ds, cfg, _ = dataset("infrastructure_health", cfg=cfg, w=w)
        doc = build(ds, cfg)
        data = render_pdf(doc, compress=False)
        self.assertIn(b"INCOMPLETE OR DEGRADED DATA", data)

    def test_truncated_table_says_so(self):
        w = world()
        for i in range(60):
            w.add_host(str(300 + i), "EXTRA-%02d" % i, "S")
        ds, cfg, _ = dataset("daily_network_health", w=w)
        doc = build(ds, cfg)
        text = pdf_text(render_pdf(doc, compress=False))
        self.assertIn(b"Showing 30 of 65 rows", text)

    def test_deterministic_bytes(self):
        doc, ds = get_docs()["executive_summary"]
        self.assertEqual(render_pdf(doc), render_pdf(doc))

    def test_compressed_pdf_is_valid_and_smaller(self):
        doc, ds = get_docs()["monitoring_quality"]
        self.assertLess(len(render_pdf(doc, compress=True)), len(render_pdf(doc, compress=False)))

    def test_non_latin_names_do_not_crash(self):
        w = world()
        w.hosts["2"]["name"] = "Router 路由器 éü"
        ds, cfg, _ = dataset("daily_network_health", w=w)
        doc = build(ds, cfg)
        self.assertTrue(render_pdf(doc).startswith(b"%PDF"))
        self.assertTrue(render_xlsx(doc))

    def test_special_characters_are_escaped(self):
        w = world()
        w.hosts["2"]["name"] = "R&D <b>core</b> (A)"
        ds, cfg, _ = dataset("daily_network_health", w=w)
        self.assertTrue(render_pdf(build(ds, cfg)).startswith(b"%PDF"))

    def test_empty_report_renders(self):
        w = world(events=False, trends=False, wan=False)
        for key in ("daily_network_health", "incident_analysis", "monitoring_quality"):
            ds, cfg, _ = dataset(key, w=w)
            doc = build(ds, cfg)
            self.assertTrue(render_pdf(doc).startswith(b"%PDF"))
            self.assertTrue(render_xlsx(doc))
            self.assertTrue(render_json(doc))


if __name__ == "__main__":
    unittest.main()
